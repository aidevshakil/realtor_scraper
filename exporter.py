import re
from pathlib import Path

from geopy.geocoders import Nominatim
from openpyxl import Workbook, load_workbook

COLUMNS = [
    "URL",
    "ADDRESS",
    "CITY",
    "PROVINCE",
    "POSTAL",
    "AGENT",
    "BROKER",
    "PRICE",
    "LATITUDE",
    "LONGITUDE",
    "MLS NUMBER",
    "BEDROOMS",
    "BATHROOMS",
    "SQFT",
    "FLOORING",
    "COOLING",
    "HEATING TYPE",
    "POOL TYPE",
    "PARKING TYPE",
    "TOTAL PARKING SPACES",
    "PROPERTY TYPE",
    "BUILDING TYPE",
    "STOREYS",
    "SQUARE FOOTAGE",
    "COMMUNITY NAME",
    "TITLE",
    "LAND SIZE",
    "AGE OF BUILDING",
    "ANNUAL PROPERTY TAXES",
    "TIME ON REALTOR",
]

DEFAULT_PATH = Path(__file__).resolve().parent / "output" / "listings.xlsx"

_geolocator = Nominatim(user_agent="realtor_scraper_personal_house_hunting")


def _clean_for_geocoding(address: str, city: str) -> tuple:
    # Drop unit/suite prefixes like "206 - 1960 QUEEN STREET" -> "1960 QUEEN STREET"
    address = re.sub(r"^\s*[\w-]+\s*-\s*", "", address or "")
    # Drop neighbourhood in parentheses like "TORONTO (THE BEACHES)" -> "TORONTO"
    city = re.sub(r"\s*\([^)]*\)", "", city or "").strip()
    return address, city


def geocode_address(address: str, city: str, province: str, postal: str) -> tuple:
    address, city = _clean_for_geocoding(address, city)
    candidates = [
        ", ".join(part for part in [address, city, province, postal, "Canada"] if part),
        ", ".join(part for part in [postal, "Canada"] if part),
    ]
    for query in candidates:
        if not query:
            continue
        try:
            location = _geolocator.geocode(query, timeout=10)
        except Exception:
            continue
        if location is not None:
            return str(location.latitude), str(location.longitude)
    return "", ""


def ensure_workbook(path: Path = DEFAULT_PATH) -> Path:
    path.parent.mkdir(exist_ok=True)
    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Listings"
        ws.append(COLUMNS)
        wb.save(path)
    return path


def _sort_by_postal(ws) -> None:
    postal_idx = COLUMNS.index("POSTAL")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    rows.sort(key=lambda r: (r[postal_idx] or ""))
    for i, row in enumerate(rows, start=2):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)


def append_listing(row: dict, path: Path = DEFAULT_PATH) -> Path:
    ensure_workbook(path)

    row = {k: (v.upper() if isinstance(v, str) else v) for k, v in row.items()}

    if not row.get("LATITUDE") and not row.get("LONGITUDE"):
        lat, lon = geocode_address(
            row.get("ADDRESS", ""),
            row.get("CITY", ""),
            row.get("PROVINCE", ""),
            row.get("POSTAL", ""),
        )
        row["LATITUDE"], row["LONGITUDE"] = lat, lon

    wb = load_workbook(path)
    ws = wb["Listings"]
    ws.append([row.get(col, "") for col in COLUMNS])
    _sort_by_postal(ws)
    wb.save(path)
    return path


if __name__ == "__main__":
    ensure_workbook()
    print("Workbook ready at:", DEFAULT_PATH)
